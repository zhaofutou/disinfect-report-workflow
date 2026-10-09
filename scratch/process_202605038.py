import docx
import sqlite3
import re
from pathlib import Path
from copy import deepcopy

docx_path = r'C:\Users\zhaoj\disinfect-report-workflow\scratch\202605038戴庄_converted.docx'
db_path = r'C:\Users\zhaoj\disinfect-report-workflow\disinfect.db'
tpl_path = r'C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\Report-templates\Template-hospital-survailance4.0.docx'
output_dir = r'C:\Users\zhaoj\disinfect-report-workflow\output'
Path(output_dir).mkdir(parents=True, exist_ok=True)
output_path = Path(output_dir) / '202605038_消毒与灭菌检验原始记录.docx'

# 1. Connect DB and insert event 202605038 records
conn = sqlite3.connect(db_path)
cur = conn.cursor()

cur.execute("DELETE FROM sample_details WHERE event = '202605038'")
conn.commit()

doc = docx.Document(docx_path)
client_name = '山东省戴庄医院'
event_id = '202605038'

records = []
for t in doc.tables:
    for row in t.rows[1:]:
        c = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        sid = c[0]
        if not sid or not ('2026' in sid or '2025' in sid):
            continue
        
        loc = c[1]
        cat = c[2]
        name = c[3]
        dur = c[4]
        area = c[5]
        bac = 1 if '√' in c[6] else 0
        rl = 1 if '√' in c[7] else 0
        jp = 1 if '√' in c[8] else 0
        tl = 1 if '√' in c[9] else 0
        uv = 1 if (c[10] and any(ch.isdigit() for ch in c[10])) else (1 if '√' in c[10] else 0)
        
        records.append((
            event_id, client_name, sid, loc, cat, name, dur, area,
            bac, rl, jp, tl, uv
        ))

cur.executemany("""
INSERT INTO sample_details 
(event, client, sampleID, sampleLocation, sampleCategory, sampleName, sampleDuration, sampleArea,
 bacCount, ronglian, jinpu, tonglv, UV)
VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
""", records)
conn.commit()
print(f"Inserted {len(records)} records for event {event_id}.")

# 2. Map observations from 202605038.jpg
obs_map = {
    # 1 - 30
    1: 2, 2: 0, 3: 43, 4: 352, 5: 65, 6: 0, 7: 0, 8: 7, 9: 0, 10: 52,
    11: 25, 12: 0, 13: 1, 14: 0, 15: 12, 16: 0, 17: 36, 18: 0, 19: 451, 20: 0,
    21: 1, 22: 0, 23: 1, 24: 1, 25: 11, 26: 5, 27: 2, 28: 646, 29: 285, 30: 0,
    # 31 - 38
    31: 0, 32: 0, 33: 0, 34: 3, 35: 0, 36: 8, 37: 0, 38: 0,
    # 39 - 44 (Air samples)
    39: "0, 6, 1",            # 三病区
    40: "1, 1, 1",            # 二病区
    41: "2, 1, 0, 1, 0",      # Mect
    42: "3, 0, 0",            # 合并症科
    43: "2, 4, 4",            # 心理6病区
    44: "1, 0, 0, 0, 0"       # 检验科
}

LIQUID_DISINFECTANTS = ['酒精', '消毒剂', '碘伏', '邻苯二甲醛', '过氧乙酸', '手消', '手消毒液']

def is_disinfectant(name):
    return any(p in name for p in LIQUID_DISINFECTANTS)

def is_water(name):
    return any(k in name for k in ['水', '漱口水', '手机水', '三枪水', '水源水', '纯化水', '漂洗水']) and not is_disinfectant(name)

def is_air(name):
    return '空气' in name or name.endswith('气')

# 3. Update database observation & result
cur.execute("SELECT sampleID, sampleName, sampleArea, sampleDuration FROM sample_details WHERE event = ? ORDER BY sampleID", (event_id,))
rows = cur.fetchall()

formatted_samples = []
for sid, name, area, dur in rows:
    num = int(sid[-3:])
    val = obs_map.get(num, 0)
    
    name_clean = name.replace('φ', '⌀').replace('Φ', '⌀')
    
    if '紫外线' in name_clean:
        obs_text = str(val) if val is not None else ""
        res_text = ""
    elif is_disinfectant(name_clean):
        m_num = re.search(r'(\d+)', str(val))
        obs_num = int(m_num.group(1)) if m_num else 0
        obs_text = f"{obs_num}(十倍稀释)"
        if obs_num == 0:
            res_text = "<10 CFU/mL"
        else:
            res_text = f"{obs_num * 10} CFU/mL"
    elif is_water(name_clean):
        obs_text = str(val)
        try:
            obs_num = float(val)
            if obs_num == 0:
                res_text = "<1.0 CFU/mL"
            else:
                res_text = f"{obs_num:.1f} CFU/mL"
        except:
            res_text = str(val)
    elif is_air(name_clean):
        obs_text = str(val)
        dur_str = f"·{dur}min" if dur else ""
        if ',' in str(val):
            plates = [float(x.strip()) for x in str(val).split(',')]
            avg = sum(plates) / len(plates)
            if avg == 0:
                thresh = 1.0 / len(plates)
                res_text = f"<{thresh:.1f} CFU/⌀90mm平皿{dur_str}"
            else:
                res_text = f"{avg:.1f} CFU/⌀90mm平皿{dur_str}"
        else:
            try:
                obs_num = float(val)
                if obs_num == 0:
                    res_text = f"0.0 CFU/⌀90mm平皿{dur_str}"
                else:
                    res_text = f"{obs_num:.1f} CFU/⌀90mm平皿{dur_str}"
            except:
                res_text = str(val)
    else:
        obs_text = str(val)
        try:
            obs_num = float(val)
            if area:
                m = re.search(r'(\d+)', str(area))
                if m:
                    area_num = float(m.group(1))
                    if obs_num == 0:
                        thresh = 10.0 / area_num
                        res_text = f"<{thresh:.1f}"
                    else:
                        res_text = f"{obs_num * 10 / area_num:.1f}"
                else:
                    res_text = str(obs_num)
            else:
                res_text = str(obs_num)
        except:
            res_text = str(val)

    cur.execute("""
    UPDATE sample_details
    SET sampleName = ?, observation = ?, result = ?
    WHERE sampleID = ?
    """, (name_clean, obs_text, res_text, sid))
    
    # Pathogen fields
    if '紫外线' in name_clean or is_air(name_clean):
        sa_val = ""
        tl_val = ""
        sh_val = ""
    else:
        sa_val = "未检出"
        tl_val = "未检出"
        sh_val = "未检出"
        
    cur.execute("SELECT sampleLocation FROM sample_details WHERE sampleID = ?", (sid,))
    loc = cur.fetchone()[0]
    
    formatted_samples.append({
        'sid': sid,
        'loc': loc,
        'name': name_clean,
        'obs': obs_text,
        'res': res_text,
        'sa': sa_val,
        'tl': tl_val,
        'sh': sh_val
    })

conn.commit()
conn.close()

# 4. Generate Word Report
doc = docx.Document(tpl_path)

def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = str(text) if text is not None else ""
        return
    p = cell.paragraphs[0]
    if not p.runs:
        p.text = str(text) if text is not None else ""
        return
    p.runs[0].text = str(text) if text is not None else ""
    for r in p.runs[1:]:
        r.text = ""

total_samples = len(formatted_samples)
entries_per_page = 18
total_pages = (total_samples + entries_per_page - 1) // entries_per_page

base_tbl_elem = doc.tables[0]._element
parent_elem = base_tbl_elem.getparent()

for t in doc.tables[1:]:
    t._element.getparent().remove(t._element)

for i in range(1, total_pages):
    new_elem = deepcopy(base_tbl_elem)
    parent_elem.append(new_elem)

event_id_full = f"202605038(202605038001-202605038{total_samples:03d})"
date_rec = "2026年9月15日"
date_fin = "2026年9月19日"

for page_idx, table in enumerate(doc.tables):
    page_num = page_idx + 1
    
    set_cell_text(table.rows[2].cells[3], event_id_full)
    set_cell_text(table.rows[2].cells[19], str(total_pages))
    set_cell_text(table.rows[2].cells[24], f"第 {page_num} 页")
    
    set_cell_text(table.rows[3].cells[3], client_name)
    set_cell_text(table.rows[3].cells[14], date_rec)
    
    set_cell_text(table.rows[4].cells[3], "细菌菌落总数、金黄色葡萄球菌、溶血性链球菌、铜绿假单胞菌")
    set_cell_text(table.rows[4].cells[14], date_fin)
    
    set_cell_text(table.rows[5].cells[3], "GB15982-2012《医院消毒卫生标准》 卫生部《消毒技术规范》（2002）")
    
    set_cell_text(table.rows[26].cells[5], f"{date_rec}15时")
    set_cell_text(table.rows[26].cells[11], f"{date_fin}10时")
    
    start_sample_idx = page_idx * entries_per_page
    for row_offset in range(entries_per_page):
        r_idx = 8 + row_offset
        sample_idx = start_sample_idx + row_offset
        
        row = table.rows[r_idx]
        
        if sample_idx < total_samples:
            item = formatted_samples[sample_idx]
            set_cell_text(row.cells[0], item['sid'])
            set_cell_text(row.cells[4], item['loc'])
            set_cell_text(row.cells[6], item['name'])
            set_cell_text(row.cells[8], item['obs'])
            set_cell_text(row.cells[13], item['res'])
            set_cell_text(row.cells[18], item['sa'])
            set_cell_text(row.cells[23], item['tl'])
            set_cell_text(row.cells[28], item['sh'])
        else:
            set_cell_text(row.cells[0], "")
            set_cell_text(row.cells[4], "")
            set_cell_text(row.cells[6], "")
            set_cell_text(row.cells[8], "")
            set_cell_text(row.cells[13], "")
            set_cell_text(row.cells[18], "")
            set_cell_text(row.cells[23], "")
            set_cell_text(row.cells[28], "")

doc.save(str(output_path))
print(f"Successfully generated 202605038 report: {output_path}")
