import docx
import sqlite3
import re
from pathlib import Path
from copy import deepcopy

db_path = r"C:\Users\zhaoj\disinfect-report-workflow\disinfect.db"
tpl_path = r"C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\Report-templates\Template-hospital-survailance4.0.docx"
output_finished_dir = r"C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\finished_report"
output_workflow_dir = r"C:\Users\zhaoj\disinfect-report-workflow\output"

target_filename = "202605032公共卫生中心洁净室.docx"
target_finished_path = Path(output_finished_dir) / target_filename
target_workflow_path = Path(output_workflow_dir) / target_filename

# 1. Update Database (disinfect.db)
conn = sqlite3.connect(db_path)
cur = conn.cursor()

# Ensure event_inspections table exists
cur.execute("""
CREATE TABLE IF NOT EXISTS event_inspections (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    eventID TEXT NOT NULL,
    sampleCount TEXT,
    clientInfo TEXT,
    inspectionDate DATE,
    inspectionTarget TEXT,
    inspectionStandards TEXT,
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);
""")

cur.execute("DELETE FROM event_inspections WHERE eventID = '202605032'")
cur.execute("DELETE FROM sample_details WHERE event = '202605032'")

# Insert event record
cur.execute("""
INSERT INTO event_inspections 
(eventID, sampleCount, clientInfo, inspectionDate, inspectionTarget, inspectionStandards)
VALUES (?, ?, ?, ?, ?, ?)
""", (
    "202605032",
    "202605032001-202605032017",
    "济宁市公共卫生医疗中心",
    "2026-07-29",
    "风速、静压差、尘埃粒子、温湿度、照度、噪声、沉降菌共7项",
    "GB50333-2013 医院洁净手术部建筑技术规范 《消毒技术规范》（2002版） GB50346-2011 生物安全实验室建筑技术规范"
))

samples_raw = [
    # sid, loc, name, plates
    ("202605032001", "手术1室", "1号手术室手术区空气", "0,0,0"),
    ("202605032002", "手术1室", "1号手术室周边区空气", "0,0,0,0,0,0"),
    ("202605032003", "手术2室", "2号手术室手术区空气", "0,0,0"),
    ("202605032004", "手术2室", "2号手术室周边区空气", "0,0,1,0,0,0"),
    ("202605032005", "手术3室", "3号手术室手术区空气", "0,0,0,0,0,0,0,0,0,0,0,2,0"),
    ("202605032006", "手术3室", "3号手术室周边区空气", "0,0,0,0,0,0,0,1"),
    ("202605032007", "手术6室", "6号手术室手术区空气", "0,0,0,0,0,0,0,0,0,0,0,1,0"),
    ("202605032008", "手术6室", "6号手术室周边区空气", "0,0,1,0,0,0,0,0"),
    ("202605032009", "手术7室", "7号手术室手术区空气", "0,0,0"),
    ("202605032010", "手术7室", "7号手术室周边区空气", "0,0,1,0,2,0"),
    ("202605032011", "手术8室", "8号手术室手术区空气", "0,0,0"),
    ("202605032012", "手术8室", "8号手术室周边区空气", "0,0,0,0,0,0"),
    ("202605032013", "手术9室", "9号手术室手术区空气", "1,0,0"),
    ("202605032014", "手术9室", "9号手术室周边区空气", "0,0,0,1,0,0"),
    ("202605032015", "PCR室", "试剂准备区空气", "0,0,0"),
    ("202605032016", "PCR室", "样本制备区空气", "0,1,1,0,1"),
    ("202605032017", "PCR室", "扩增分析区空气", "3,5,0"),
]

formatted_samples = []
for sid, loc, name, obs in samples_raw:
    plates = [float(x.strip()) for x in obs.split(',')]
    n = len(plates)
    s = sum(plates)
    if s == 0:
        thresh = 1.0 / n
        res = f"<{thresh:.1f} CFU/⌀90mm皿·30min"
    else:
        avg = s / n
        res = f"{avg:.1f} CFU/⌀90mm皿·30min"
    
    formatted_samples.append({
        'sid': sid,
        'loc': loc,
        'name': name,
        'obs': obs,
        'res': res
    })
    
    cur.execute("""
    INSERT INTO sample_details
    (event, client, sampleID, sampleLocation, sampleCategory, sampleName, sampleDuration, observation, result)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, ("202605032", "济宁市公共卫生医疗中心", sid, loc, "I", name, "30", obs, res))

conn.commit()
conn.close()
print(f"✓ Database updated for 202605032 ({len(formatted_samples)} records).")

# 2. Generate Word Report
doc = docx.Document(tpl_path)

def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = str(text) if text is not None else ""
        return
    p = cell.paragraphs[0]
    if not p.runs:
        p.text = str(text) if text is not None else ""
        return
    p.runs[0].text = str(text) if text is not None else ""
    for r in p.runs[1:]:
        r.text = ""

# Since total samples = 17 <= 18, exactly 1 page (1 table)
# Remove extra tables from template
for t in doc.tables[1:]:
    t._element.getparent().remove(t._element)

table = doc.tables[0]
event_id_full = "202605032(202605032001-202605032017)"
client_name = "济宁市公共卫生医疗中心"
date_rec = "2026年7月29日"
date_fin = "2026年8月1日"
standards = "GB50333-2013《医院洁净手术部建筑技术规范》 卫生部《消毒技术规范》（2002版） GB50346-2011《生物安全实验室建筑技术规范》"

# Header metadata
set_cell_text(table.rows[2].cells[3], event_id_full)
set_cell_text(table.rows[2].cells[19], "1")
set_cell_text(table.rows[2].cells[24], "第 1 页")

set_cell_text(table.rows[3].cells[3], client_name)
set_cell_text(table.rows[3].cells[14], date_rec)

set_cell_text(table.rows[4].cells[3], "细菌菌落总数")
set_cell_text(table.rows[4].cells[14], date_fin)

set_cell_text(table.rows[5].cells[3], standards)

set_cell_text(table.rows[26].cells[5], f"{date_rec}15时")
set_cell_text(table.rows[26].cells[11], f"{date_fin}16时")

# Populate 18 rows (rows 8..25)
for offset in range(18):
    r_idx = 8 + offset
    row = table.rows[r_idx]
    if offset < len(formatted_samples):
        item = formatted_samples[offset]
        set_cell_text(row.cells[0], item['sid'])
        set_cell_text(row.cells[4], item['loc'])
        set_cell_text(row.cells[6], item['name'])
        set_cell_text(row.cells[8], item['obs'])
        set_cell_text(row.cells[13], item['res'])
        # Pathogen columns left blank
        set_cell_text(row.cells[18], "")
        set_cell_text(row.cells[23], "")
        set_cell_text(row.cells[28], "")
    else:
        # Blank padding row
        set_cell_text(row.cells[0], "")
        set_cell_text(row.cells[4], "")
        set_cell_text(row.cells[6], "")
        set_cell_text(row.cells[8], "")
        set_cell_text(row.cells[13], "")
        set_cell_text(row.cells[18], "")
        set_cell_text(row.cells[23], "")
        set_cell_text(row.cells[28], "")

doc.save(str(target_finished_path))
doc.save(str(target_workflow_path))

print(f"✓ Report successfully generated at: {target_finished_path}")
print(f"✓ Backup copy saved at: {target_workflow_path}")
