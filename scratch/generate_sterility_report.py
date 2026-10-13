import docx
import sqlite3
import shutil
from pathlib import Path

# Paths
base_template_path = r"C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\finished_report\202505056公卫中心无菌器械.docx"
output_finished_dir = r"C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\finished_report"
output_workflow_dir = r"C:\Users\zhaoj\disinfect-report-workflow\output"
db_path = r"C:\Users\zhaoj\disinfect-report-workflow\disinfect.db"

target_filename = "202605030公卫中心无菌器械.docx"
target_finished_path = Path(output_finished_dir) / target_filename
target_workflow_path = Path(output_workflow_dir) / target_filename

# 1. Update Database (disinfect.db)
conn = sqlite3.connect(db_path)
cur = conn.cursor()

# Ensure event_inspections table exists
cur.execute("""
CREATE TABLE IF NOT EXISTS event_inspections (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    eventID TEXT NOT NULL,
    sampleCount TEXT,
    clientInfo TEXT,
    inspectionDate DATE,
    inspectionTarget TEXT,
    inspectionStandards TEXT,
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);
""")

# Delete existing records for event 202605030 if any
cur.execute("DELETE FROM event_inspections WHERE eventID = '202605030'")
cur.execute("DELETE FROM sample_details WHERE event = '202605030'")

# Insert event record
cur.execute("""
INSERT INTO event_inspections 
(eventID, sampleCount, clientInfo, inspectionDate, inspectionTarget, inspectionStandards)
VALUES (?, ?, ?, ?, ?, ?)
""", (
    "202605030",
    "202605030001-202605030002",
    "济宁市公共卫生医疗中心",
    "2026-08-03",
    "无菌试验（需氧-厌氧菌，霉菌）",
    "GB15982-2012《医院消毒卫生标准》 卫生部《消毒技术规范》（2002年版）"
))

# Insert sample details
samples = [
    ("202605030", "济宁市公共卫生医疗中心", "202605030001", "结核科", "无菌器械", "转换头1", "-", "无菌生长"),
    ("202605030", "济宁市公共卫生医疗中心", "202605030002", "结核科", "无菌器械", "转换头2", "-", "无菌生长"),
]

cur.executemany("""
INSERT INTO sample_details
(event, client, sampleID, sampleLocation, sampleCategory, sampleName, observation, result)
VALUES (?, ?, ?, ?, ?, ?, ?, ?)
""", samples)

conn.commit()
print("✓ Database updated successfully (event_inspections & sample_details)")

# 2. Generate Word Report
doc = docx.Document(base_template_path)
t = doc.tables[0]

def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = str(text) if text is not None else ""
        return
    p = cell.paragraphs[0]
    if not p.runs:
        p.text = str(text) if text is not None else ""
        return
    p.runs[0].text = str(text) if text is not None else ""
    for r in p.runs[1:]:
        r.text = ""

# Row 0: Sample ID
set_cell_text(t.rows[0].cells[12], "202605030(202605030001-202605030002)")

# Row 1: Sample Name & Sterilization/Inspection Date
set_cell_text(t.rows[1].cells[2], "无菌器械(结核科)")
set_cell_text(t.rows[1].cells[12], "2026年8月3日")

# Row 7 & 8: Item 1 (转换头1)
set_cell_text(t.rows[7].cells[0], "202605030001转换头1")
set_cell_text(t.rows[8].cells[0], "202605030001转换头1")

# Row 9 & 10: Item 2 (转换头2)
set_cell_text(t.rows[9].cells[0], "202605030002转换头2")
set_cell_text(t.rows[10].cells[0], "202605030002转换头2")

# Row 19: Inspection Date Range (8月3日 - 8月10日)
set_cell_text(t.rows[19].cells[2], "2026年8月3日15时")
set_cell_text(t.rows[19].cells[8], "2026年8月10日9时")

# Save outputs
doc.save(str(target_finished_path))
doc.save(str(target_workflow_path))

print(f"✓ Report generated at: {target_finished_path}")
print(f"✓ Copy saved at: {target_workflow_path}")
