import docx
import sqlite3
import re
from pathlib import Path
from copy import deepcopy

db_path = r'C:\Users\zhaoj\disinfect-report-workflow\disinfect.db'
tpl_path = r'C:\Users\zhaoj\OneDrive\2026工作\消毒报告3.0版\Report-templates\Template-hospital-survailance4.0.docx'
output_dir = r'C:\Users\zhaoj\disinfect-report-workflow\output'
Path(output_dir).mkdir(parents=True, exist_ok=True)
output_path = Path(output_dir) / '202605029_消毒与灭菌检验原始记录.docx'

# 1. Query records from SQLite DB
conn = sqlite3.connect(db_path)
cur = conn.cursor()

cur.execute("""
SELECT sampleID, sampleLocation, sampleName, sampleArea, sampleDuration, observation, result, jinpu, ronglian, tonglv, UV
FROM sample_details
ORDER BY id
""")
db_rows = cur.fetchall()
conn.close()

total_samples = len(db_rows)
print(f"Loaded {total_samples} samples from DB.")

# 2. Setup formatting helper
def set_cell_text(cell, text):
    if not cell.paragraphs:
        cell.text = str(text) if text is not None else ""
        return
    p = cell.paragraphs[0]
    if not p.runs:
        p.text = str(text) if text is not None else ""
        return
    p.runs[0].text = str(text) if text is not None else ""
    for r in p.runs[1:]:
        r.text = ""

LIQUID_DISINFECTANTS = ['酒精', '消毒剂', '碘伏', '邻苯二甲醛', '过氧乙酸', '手消']

def is_disinfectant(name):
    return any(p in name for p in LIQUID_DISINFECTANTS)

def is_endoscope(name):
    return any(k in name for k in ['胃镜', '肠镜', '气管镜', '支气管镜', '喉镜'])

def is_water(name):
    return any(k in name for k in ['漂洗水', '纯化水', '透析水']) and not is_endoscope(name) and not is_disinfectant(name)

def is_air(name):
    return ('空气' in name or name.endswith('气')) and not is_endoscope(name)

# 3. Format row data according to guidelines
formatted_samples = []
for r in db_rows:
    sid, loc, name, area, dur, obs, res, jp, rl, tl, uv = r
    
    # Character replacement φ -> ⌀
    name_clean = name.replace('φ', '⌀').replace('Φ', '⌀')
    
    # Observation formatting
    if '紫外线' in name_clean:
        obs_val = str(obs) if obs is not None else ""
        res_val = ""
    elif is_disinfectant(name_clean):
        m_num = re.search(r'(\d+)', str(obs))
        obs_num = int(m_num.group(1)) if m_num else 0
        obs_val = f"{obs_num}(十倍稀释)"
        if obs_num == 0:
            res_val = "<10 CFU/mL"
        else:
            res_val = f"{obs_num * 10} CFU/mL"
    else:
        obs_val = str(obs) if obs is not None else ""
        res_val = str(res) if res is not None else ""
    
    # Result character replacement φ -> ⌀
    if res_val:
        res_val = res_val.replace('φ', '⌀').replace('Φ', '⌀')
        
    # Pathogens formatting
    if '紫外线' in name_clean or is_air(name_clean):
        sa_val = ""
        tl_val = ""
        sh_val = ""
    else:
        sa_val = "未检出" if jp == 0 else "检出"
        tl_val = "未检出" if tl == 0 else "检出"
        sh_val = "未检出" if rl == 0 else "检出"
        
    formatted_samples.append({
        'sid': sid,
        'loc': loc,
        'name': name_clean,
        'obs': obs_val,
        'res': res_val,
        'sa': sa_val,
        'tl': tl_val,
        'sh': sh_val
    })

# 4. Load Word Template and Expand
doc = docx.Document(tpl_path)

entries_per_page = 18
total_pages = (total_samples + entries_per_page - 1) // entries_per_page

print(f"Total pages required: {total_pages}")

base_tbl_elem = doc.tables[0]._element
parent_elem = base_tbl_elem.getparent()

# Remove initial extra tables
for t in doc.tables[1:]:
    t._element.getparent().remove(t._element)

# Clone base table for total_pages
for i in range(1, total_pages):
    new_elem = deepcopy(base_tbl_elem)
    parent_elem.append(new_elem)

print(f"Generated {len(doc.tables)} tables in document.")

# Header info
event_id_full = f"202605029(202605029001-202605029{total_samples:03d})"
client_name = "济宁市公共卫生医疗中心"
date_rec = "2026年7月29日"
date_fin = "2026年8月2日"

# 5. Populate tables
for page_idx, table in enumerate(doc.tables):
    page_num = page_idx + 1
    
    # Metadata rows
    set_cell_text(table.rows[2].cells[3], event_id_full)
    set_cell_text(table.rows[2].cells[19], str(total_pages))
    set_cell_text(table.rows[2].cells[24], f"第 {page_num} 页")
    
    set_cell_text(table.rows[3].cells[3], client_name)
    set_cell_text(table.rows[3].cells[14], date_rec)
    
    set_cell_text(table.rows[4].cells[3], "细菌菌落总数、金黄色葡萄球菌、溶血性链球菌、铜绿假单胞菌")
    set_cell_text(table.rows[4].cells[14], date_fin)
    
    set_cell_text(table.rows[5].cells[3], "GB15982-2012《医院消毒卫生标准》 卫生部《消毒技术规范》（2002）")
    
    # Footer rows
    set_cell_text(table.rows[26].cells[5], f"{date_rec}15时")
    set_cell_text(table.rows[26].cells[11], f"{date_fin}10时")
    
    # Data rows (8..25)
    start_sample_idx = page_idx * entries_per_page
    for row_offset in range(entries_per_page):
        r_idx = 8 + row_offset
        sample_idx = start_sample_idx + row_offset
        
        row = table.rows[r_idx]
        
        if sample_idx < total_samples:
            item = formatted_samples[sample_idx]
            set_cell_text(row.cells[0], item['sid'])
            set_cell_text(row.cells[4], item['loc'])
            set_cell_text(row.cells[6], item['name'])
            set_cell_text(row.cells[8], item['obs'])
            set_cell_text(row.cells[13], item['res'])
            set_cell_text(row.cells[18], item['sa'])
            set_cell_text(row.cells[23], item['tl'])
            set_cell_text(row.cells[28], item['sh'])
        else:
            # Empty padding row
            set_cell_text(row.cells[0], "")
            set_cell_text(row.cells[4], "")
            set_cell_text(row.cells[6], "")
            set_cell_text(row.cells[8], "")
            set_cell_text(row.cells[13], "")
            set_cell_text(row.cells[18], "")
            set_cell_text(row.cells[23], "")
            set_cell_text(row.cells[28], "")

# Save document
doc.save(str(output_path))
print(f"Report generated successfully: {output_path}")
